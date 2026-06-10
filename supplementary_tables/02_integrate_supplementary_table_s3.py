# -*- coding: utf-8 -*-
"""Integrate Supplementary Table S3 into the corrected supplementary material.

The script appends S3 from the stored source document into:
论文写作/补充材料/Supplementary_Online_5_corrected_BNA_labels.docx

It preserves S1/S2 content, updates the content list, appends S3 on a landscape
page, and normalizes all text to Times New Roman.
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
SUPP_DIR = ROOT / "论文写作" / "补充材料"
MAIN = SUPP_DIR / "Supplementary_Online_5_corrected_BNA_labels.docx"
S3_SRC = SUPP_DIR / "Supplementary_Table_S3_0.3.docx"


def set_run_font(run, size_pt: float | None = None) -> None:
    run.font.name = "Times New Roman"
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), "Times New Roman")


def set_cell_margins(cell, top=55, start=50, bottom=55, end=50) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_autofit(table) -> None:
    table.autofit = True
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is not None:
        tbl_pr.remove(layout)
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")


def normalize_table(table, font_size: float = 10.0) -> None:
    set_table_autofit(table)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        for h in tr_pr.findall(qn("w:trHeight")):
            tr_pr.remove(h)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    set_run_font(run, font_size)


def normalize_doc(doc: Document) -> None:
    for style in doc.styles:
        try:
            style.font.name = "Times New Roman"
            if style._element.rPr is not None and style._element.rPr.rFonts is not None:
                for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                    style._element.rPr.rFonts.set(qn(f"w:{key}"), "Times New Roman")
        except Exception:
            pass
    for para in doc.paragraphs:
        for run in para.runs:
            set_run_font(run, 10.0 if para.text.startswith(("Note.", "Abbreviations.")) else None)
    for table in doc.tables:
        normalize_table(table, 10.0)


def paragraph_texts(doc: Document) -> list[str]:
    return [" ".join(p.text.split()) for p in doc.paragraphs]


def add_content_list_entry(doc: Document) -> None:
    entry = "Supplementary Table S3. Extended demographic, clinical, source/design, and neuroimaging quality-control characteristics of the analytic sample."
    if any(entry in p.text for p in doc.paragraphs):
        return
    # Insert after S2 content-list entry, before the repeated S1 title.
    paras = doc.paragraphs
    target_idx = None
    for idx, p in enumerate(paras):
        if p.text.strip().startswith("Supplementary Table S2. Site/study proxy"):
            target_idx = idx
            break
    if target_idx is None:
        raise RuntimeError("Could not find S2 content-list entry.")
    new_p = deepcopy(paras[target_idx]._p)
    for child in list(new_p):
        new_p.remove(child)
    paras[target_idx]._p.addnext(new_p)
    from docx.text.paragraph import Paragraph

    para = Paragraph(new_p, paras[target_idx]._parent)
    para.style = paras[target_idx].style
    run = para.add_run(entry)
    set_run_font(run, 10.0)


def set_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    # Keep the existing supplementary-landscape geometry.
    section.page_width = 10058400
    section.page_height = 7772400
    section.left_margin = 411480
    section.right_margin = 411480
    section.top_margin = 411480
    section.bottom_margin = 411480


def append_s3(main: Document, s3: Document) -> None:
    texts = paragraph_texts(s3)
    s3_title = next(t for t in texts if t.startswith("Supplementary Table S3."))
    panel_titles = [t for t in texts if t.startswith("Panel ")]
    notes = [t for t in texts if t.startswith("Note.")]
    abbreviations = [t for t in texts if t.startswith("Abbreviations.")]
    if len(panel_titles) != len(s3.tables):
        raise RuntimeError(f"Panel/table mismatch: {len(panel_titles)} panels, {len(s3.tables)} tables.")

    main.add_section(WD_SECTION.NEW_PAGE)
    set_landscape(main.sections[-1])

    p = main.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(s3_title)
    r.bold = True
    set_run_font(r, 10.0)

    for panel, table in zip(panel_titles, s3.tables):
        p = main.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(panel)
        r.bold = True
        set_run_font(r, 10.0)
        new_tbl = deepcopy(table._tbl)
        main._body._element.append(new_tbl)
        normalize_table(main.tables[-1], 10.0)

    for text in notes + abbreviations:
        p = main.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        set_run_font(r, 10.0)


def main() -> None:
    doc = Document(str(MAIN))
    s3 = Document(str(S3_SRC))
    if any("Supplementary Table S3. Extended demographic" in p.text for p in doc.paragraphs):
        raise RuntimeError("S3 already appears in main supplementary document; refusing to duplicate.")
    add_content_list_entry(doc)
    append_s3(doc, s3)
    normalize_doc(doc)
    doc.save(str(MAIN))
    print(f"Updated {MAIN}")


if __name__ == "__main__":
    main()
